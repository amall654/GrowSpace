from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
d=Document()
for a in ['left','right','start','end']:
 p=d.add_paragraph();pr=p._p.get_or_add_pPr()
 b=OxmlElement('w:bidi');pr.append(b)
 j=OxmlElement('w:jc');j.set(qn('w:val'),a);pr.append(j)
 r=p.add_run(a+'  هذا اختبار لمحاذاة فقرة عربية مع نص إنجليزي GrowSpace في نفس السطر');r.font.name='Arial';r.font.size=Pt(13)
e='C:/Users/HP/Documents/project/tmp/growspace-srs/rtl-probe.docx';d.save(e)
